# Hot reload trigger
from fastapi import FastAPI, Depends, HTTPException, Query
from fastapi.responses import RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from typing import List, Optional
from datetime import datetime
from sqlmodel import Session, select
from contextlib import asynccontextmanager
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from pydantic import BaseModel

from models import Trabajo, Actividad, Tarea
from database import engine, init_db

# --- PYDANTIC SCHEMAS PARA ENTRADA ---
class TrabajoInput(BaseModel):
    titulo: str
    descripcion: Optional[str] = None
    cliente_contacto: Optional[str] = None
    estado: str = "Planificado"
    color: Optional[str] = "#3b82f6"
    nota: Optional[str] = None
    planta: Optional[str] = None
    tipo_actividad: Optional[str] = None
    importancia: Optional[str] = None
    orden: Optional[int] = 0

class ActividadInput(BaseModel):
    titulo: str
    descripcion: Optional[str] = None
    fecha_inicio: Optional[datetime] = None
    fecha_fin: Optional[datetime] = None
    estado: str = "Planificado"
    color: Optional[str] = "#3b82f6"
    nota: Optional[str] = None
    planta: Optional[str] = None
    tipo_actividad: Optional[str] = None
    importancia: Optional[str] = None
    tiempo_invertir: Optional[str] = None
    trabajo_id: Optional[int] = None
    tarea_id: Optional[int] = None
    orden: Optional[int] = 0
    duracion_minutos: Optional[int] = 0
    completado: Optional[bool] = False

class TareaInput(BaseModel):
    titulo: str
    descripcion: Optional[str] = None
    prioridad: str = "Media"
    estado: str = "Pendiente"
    trabajo_id: Optional[int] = None
    nota: Optional[str] = None
    planta: Optional[str] = None
    tipo_actividad: Optional[str] = None
    importancia: Optional[str] = None
    orden: Optional[int] = 0


# --- PYDANTIC SCHEMAS PARA RESPUESTA (CON RELACIONES) ---
class ActividadRead(BaseModel):
    id: int
    titulo: str
    descripcion: Optional[str] = None
    fecha_inicio: Optional[datetime] = None
    fecha_fin: Optional[datetime] = None
    estado: str
    color: Optional[str] = None
    orden: int
    fecha_creacion: Optional[datetime] = None
    trabajo_id: Optional[int] = None
    tarea_id: Optional[int] = None
    nota: Optional[str] = None
    planta: Optional[str] = None
    tipo_actividad: Optional[str] = None
    importancia: Optional[str] = None
    tiempo_invertir: Optional[str] = None
    duracion_minutos: Optional[int] = 0
    completado: bool

    class Config:
        from_attributes = True

class TareaRead(BaseModel):
    id: int
    titulo: str
    descripcion: Optional[str] = None
    prioridad: str
    estado: str
    orden: int
    fecha_creacion: Optional[datetime] = None
    trabajo_id: Optional[int] = None
    nota: Optional[str] = None
    planta: Optional[str] = None
    tipo_actividad: Optional[str] = None
    importancia: Optional[str] = None
    actividades: List[ActividadRead] = []

    class Config:
        from_attributes = True

class TrabajoRead(BaseModel):
    id: int
    titulo: str
    descripcion: Optional[str] = None
    cliente_contacto: Optional[str] = None
    estado: str
    color: Optional[str] = None
    orden: int
    fecha_creacion: Optional[datetime] = None
    nota: Optional[str] = None
    planta: Optional[str] = None
    tipo_actividad: Optional[str] = None
    importancia: Optional[str] = None
    actividades: List[ActividadRead] = []
    tareas: List[TareaRead] = []

    class Config:
        from_attributes = True


@asynccontextmanager
async def lifespan(app: FastAPI):
    # Inicializar Base de Datos y precargar semillas
    init_db()
    yield

app = FastAPI(
    title="Agenda de Trabajos y Tareas API",
    description="API para gestionar la agenda de trabajos, tareas y listas de pendientes en tema claro",
    lifespan=lifespan
)

# Configurar CORS para desarrollo local y remoto
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["*"],
)

# Dependencia para obtener la sesión de base de datos
def get_db():
    with Session(engine) as session:
        yield session

# Redirección de la raíz a la página de inicio
@app.get("/")
def read_root():
    return RedirectResponse(url="/static/index.html")

# --- ENDPOINTS TRABAJOS ---

@app.get("/api/trabajos", response_model=List[TrabajoRead])
def get_trabajos(db: Session = Depends(get_db)):
    stmt = select(Trabajo).order_by(Trabajo.orden)
    trabajos = db.exec(stmt).all()
    for t in trabajos:
        t.tareas.sort(key=lambda x: (x.orden, x.id))
    return trabajos

@app.get("/api/trabajos/{trabajo_id}", response_model=TrabajoRead)
def get_trabajo(trabajo_id: int, db: Session = Depends(get_db)):
    trabajo = db.get(Trabajo, trabajo_id)
    if not trabajo:
        raise HTTPException(status_code=404, detail="Trabajo no encontrado")
    trabajo.tareas.sort(key=lambda x: (x.orden, x.id))
    return trabajo

@app.post("/api/trabajos", response_model=Trabajo)
def create_trabajo(trabajo_in: TrabajoInput, db: Session = Depends(get_db)):
    # Calcular el siguiente orden
    max_orden = db.exec(select(Trabajo.orden)).all()
    next_orden = max(max_orden) + 1 if max_orden else 1
    
    data = trabajo_in.dict()
    data.pop("orden", None)
    db_trabajo = Trabajo(
        **data,
        orden=next_orden
    )
    db.add(db_trabajo)
    db.commit()
    db.refresh(db_trabajo)
    return db_trabajo

@app.put("/api/trabajos/{trabajo_id}", response_model=Trabajo)
def update_trabajo(trabajo_id: int, trabajo_data: TrabajoInput, db: Session = Depends(get_db)):
    db_trabajo = db.get(Trabajo, trabajo_id)
    if not db_trabajo:
        raise HTTPException(status_code=404, detail="Trabajo no encontrado")
    
    # Actualizar campos del modelo de base de datos desde los datos de entrada
    for key, value in trabajo_data.dict().items():
        setattr(db_trabajo, key, value)
        
    db.add(db_trabajo)
    db.commit()
    db.refresh(db_trabajo)
    return db_trabajo

@app.delete("/api/trabajos/{trabajo_id}")
def delete_trabajo(trabajo_id: int, db: Session = Depends(get_db)):
    db_trabajo = db.get(Trabajo, trabajo_id)
    if not db_trabajo:
        raise HTTPException(status_code=404, detail="Trabajo no encontrado")
    
    db.delete(db_trabajo)
    db.commit()
    return {"status": "ok", "message": f"Trabajo {trabajo_id} eliminado exitosamente."}

@app.put("/api/trabajos/reordenar")
def reorder_trabajos(ids: List[int], db: Session = Depends(get_db)):
    for idx, trabajo_id in enumerate(ids):
        db_trabajo = db.get(Trabajo, trabajo_id)
        if db_trabajo:
            db_trabajo.orden = idx + 1
            db.add(db_trabajo)
    db.commit()
    return {"status": "ok"}

# --- ENDPOINTS ACTIVIDADES ---

@app.get("/api/actividades", response_model=List[ActividadRead])
def get_actividades(
    trabajo_id: Optional[int] = None,
    tarea_id: Optional[int] = None,
    standalone: Optional[bool] = None,
    db: Session = Depends(get_db)
):
    stmt = select(Actividad)
    if trabajo_id is not None:
        stmt = stmt.where(Actividad.trabajo_id == trabajo_id)
    if tarea_id is not None:
        stmt = stmt.where(Actividad.tarea_id == tarea_id)
    if standalone is True:
        stmt = stmt.where(Actividad.trabajo_id == None).where(Actividad.tarea_id == None)
        
    stmt = stmt.order_by(Actividad.orden, Actividad.id)
    actividades = db.exec(stmt).all()
    return actividades

@app.get("/api/actividades/{actividad_id}", response_model=ActividadRead)
def get_actividad(actividad_id: int, db: Session = Depends(get_db)):
    act = db.get(Actividad, actividad_id)
    if not act:
        raise HTTPException(status_code=404, detail="Actividad no encontrada")
    return act

@app.post("/api/actividades", response_model=Actividad)
def create_actividad(act_in: ActividadInput, db: Session = Depends(get_db)):
    stmt_orden = select(Actividad.orden)
    if act_in.trabajo_id is not None:
        stmt_orden = stmt_orden.where(Actividad.trabajo_id == act_in.trabajo_id)
    elif act_in.tarea_id is not None:
        stmt_orden = stmt_orden.where(Actividad.tarea_id == act_in.tarea_id)
    else:
        stmt_orden = stmt_orden.where(Actividad.trabajo_id == None).where(Actividad.tarea_id == None)
    ordenes = db.exec(stmt_orden).all()
    next_orden = max(ordenes) + 1 if ordenes else 1
    
    data = act_in.dict()
    data.pop("orden", None)
    db_act = Actividad(
        **data,
        orden=next_orden
    )
    db.add(db_act)
    db.commit()
    db.refresh(db_act)
    return db_act

@app.put("/api/actividades/{actividad_id}", response_model=Actividad)
def update_actividad(actividad_id: int, act_data: ActividadInput, db: Session = Depends(get_db)):
    db_act = db.get(Actividad, actividad_id)
    if not db_act:
        raise HTTPException(status_code=404, detail="Actividad no encontrada")
        
    if db_act.trabajo_id != act_data.trabajo_id or db_act.tarea_id != act_data.tarea_id:
        stmt_orden = select(Actividad.orden)
        if act_data.trabajo_id is not None:
            stmt_orden = stmt_orden.where(Actividad.trabajo_id == act_data.trabajo_id)
        elif act_data.tarea_id is not None:
            stmt_orden = stmt_orden.where(Actividad.tarea_id == act_data.tarea_id)
        else:
            stmt_orden = stmt_orden.where(Actividad.trabajo_id == None).where(Actividad.tarea_id == None)
        ordenes = db.exec(stmt_orden).all()
        db_act.orden = max(ordenes) + 1 if ordenes else 1
        
    for key, value in act_data.dict().items():
        setattr(db_act, key, value)
        
    db.add(db_act)
    db.commit()
    db.refresh(db_act)
    return db_act

@app.put("/api/actividades/{actividad_id}/mover")
def move_actividad_dates(actividad_id: int, fecha_inicio: str, fecha_fin: str, db: Session = Depends(get_db)):
    db_act = db.get(Actividad, actividad_id)
    if not db_act:
        raise HTTPException(status_code=404, detail="Actividad no encontrada")
    try:
        db_act.fecha_inicio = datetime.fromisoformat(fecha_inicio.replace("Z", "+00:00"))
        db_act.fecha_fin = datetime.fromisoformat(fecha_fin.replace("Z", "+00:00"))
    except ValueError:
        raise HTTPException(status_code=400, detail="Formato de fecha inválido. Usar ISO 8601")
    db.add(db_act)
    db.commit()
    return {"status": "ok", "actividad_id": actividad_id, "fecha_inicio": db_act.fecha_inicio, "fecha_fin": db_act.fecha_fin}

@app.delete("/api/actividades/{actividad_id}")
def delete_actividad(actividad_id: int, db: Session = Depends(get_db)):
    db_act = db.get(Actividad, actividad_id)
    if not db_act:
        raise HTTPException(status_code=404, detail="Actividad no encontrada")
    db.delete(db_act)
    db.commit()
    return {"status": "ok", "message": f"Actividad {actividad_id} eliminada exitosamente."}

@app.put("/api/actividades/reordenar")
def reorder_actividades(ids: List[int], db: Session = Depends(get_db)):
    for idx, act_id in enumerate(ids):
        db_act = db.get(Actividad, act_id)
        if db_act:
            db_act.orden = idx + 1
            db.add(db_act)
    db.commit()
    return {"status": "ok"}

# --- ENDPOINTS TAREAS ---

@app.get("/api/tareas", response_model=List[TareaRead])
def get_tareas(
    trabajo_id: Optional[int] = None, 
    standalone: Optional[bool] = None,
    db: Session = Depends(get_db)
):
    stmt = select(Tarea)
    if trabajo_id is not None:
        stmt = stmt.where(Tarea.trabajo_id == trabajo_id)
    elif standalone is True:
        stmt = stmt.where(Tarea.trabajo_id == None)
        
    stmt = stmt.order_by(Tarea.orden, Tarea.id)
    tareas = db.exec(stmt).all()
    for t in tareas:
        t.actividades.sort(key=lambda x: (x.orden, x.id))
    return tareas

@app.post("/api/tareas", response_model=Tarea)
def create_tarea(tarea_in: TareaInput, db: Session = Depends(get_db)):
    stmt_orden = select(Tarea.orden)
    if tarea_in.trabajo_id is not None:
        stmt_orden = stmt_orden.where(Tarea.trabajo_id == tarea_in.trabajo_id)
    else:
        stmt_orden = stmt_orden.where(Tarea.trabajo_id == None)
    ordenes = db.exec(stmt_orden).all()
    next_orden = max(ordenes) + 1 if ordenes else 1
    
    data = tarea_in.dict()
    data.pop("orden", None)
    db_tarea = Tarea(
        **data,
        orden=next_orden
    )
    
    db.add(db_tarea)
    db.commit()
    db.refresh(db_tarea)
    return db_tarea

@app.put("/api/tareas/{tarea_id}", response_model=Tarea)
def update_tarea(tarea_id: int, tarea_data: TareaInput, db: Session = Depends(get_db)):
    db_tarea = db.get(Tarea, tarea_id)
    if not db_tarea:
        raise HTTPException(status_code=404, detail="Tarea no encontrada")
    
    # Si cambió de trabajo_id, recalculamos el orden en el nuevo destino
    if db_tarea.trabajo_id != tarea_data.trabajo_id:
        stmt_orden = select(Tarea.orden)
        if tarea_data.trabajo_id is not None:
            stmt_orden = stmt_orden.where(Tarea.trabajo_id == tarea_data.trabajo_id)
        else:
            stmt_orden = stmt_orden.where(Tarea.trabajo_id == None)
        ordenes = db.exec(stmt_orden).all()
        db_tarea.orden = max(ordenes) + 1 if ordenes else 1
    
    for key, value in tarea_data.dict().items():
        setattr(db_tarea, key, value)
        
    db.add(db_tarea)
    db.commit()
    db.refresh(db_tarea)
    return db_tarea

@app.delete("/api/tareas/{tarea_id}")
def delete_tarea(tarea_id: int, db: Session = Depends(get_db)):
    db_tarea = db.get(Tarea, tarea_id)
    if not db_tarea:
        raise HTTPException(status_code=404, detail="Tarea no encontrada")
    
    db.delete(db_tarea)
    db.commit()
    return {"status": "ok", "message": f"Tarea {tarea_id} eliminada exitosamente."}

@app.put("/api/tareas/reordenar")
def reorder_tareas(ids: List[int], db: Session = Depends(get_db)):
    for idx, tarea_id in enumerate(ids):
        db_tarea = db.get(Tarea, tarea_id)
        if db_tarea:
            db_tarea.orden = idx + 1
            db.add(db_tarea)
    db.commit()
    return {"status": "ok"}

# --- EXPORTACIÓN EXCEL ---

@app.get("/api/exportar/excel")
def exportar_excel(db: Session = Depends(get_db)):
    wb = openpyxl.Workbook()
    
    font_bold = Font(name="Outfit", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=11, bold=False)
    
    align_header_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_header_left = Alignment(horizontal="left", vertical="center")
    
    align_data_center = Alignment(horizontal="center", vertical="center")
    align_data_left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_data_left = Alignment(horizontal="left", vertical="center")
    
    widths_trabajos = {
        "ID Trabajo": 6.36,
        "Título Trabajo": 31.9,
        "Descripción": 59.45,
        "Contacto/Cliente": 19.0,
        "Estado": 11.18,
        "Color": 7.9,
        "Fecha Creación": 14.18,
        "Nota": 52.73,
        "Planta": 8.54,
        "Tipo Actividad": 14.54,
        "Importancia": 12.82
    }
    widths_gestiones = {
        "ID Trabajo": 7.82,
        "Título Trabajo": 23.73,
        "ID Gestión": 6.18,
        "Título Gestión": 43.09,
        "Descripción": 51.0,
        "Prioridad": 12.0,
        "Estado": 11.63,
        "Fecha Creación": 14.36,
        "Nota": 28.18,
        "Planta": 12.0,
        "Tipo Actividad": 8.45,
        "Importancia": 14.0
    }
    widths_actividades = {
        "ID Actividad": 6.54,
        "Título Actividad": 22.0,
        "Descripción": 33.0,
        "Fecha Inicio": 13.09,
        "Fecha Fin": 13.0,
        "Estado": 11.27,
        "Color": 7.45,
        "ID Trabajo": 6.82,
        "Título Trabajo": 24.54,
        "ID Gestión": 6.82,
        "Título Gestión": 24.54,
        "Fecha Creación": 17.0,
        "Nota": 48.18,
        "Planta": 12.0,
        "Tipo Actividad": 14.82,
        "Importancia": 13.82,
        "Tiempo a Invertir": 17.73,
        "Completado": 12.0
    }
    widths_todos = {
        "ID Actividad": 12.0,
        "Título Actividad": 39.0,
        "Completado": 13.0,
        "Fecha Creación": 17.0,
        "Descripción": 39.0,
        "Planta": 13.0,
        "Importancia": 15.0
    }
    widths_consolidado = {
        "ID Trabajo": 6.27,
        "Título Trabajo": 25.0,
        "Estado Trabajo": 12.54,
        "Contacto/Cliente": 16.45,
        "ID Gestión": 7.36,
        "Título Gestión": 25.0,
        "Estado Gestión": 12.27,
        "Prioridad Gestión": 14.45,
        "ID Actividad": 12.0,
        "Título Actividad": 23.0,
        "Descripción Actividad": 35.0,
        "Estado Actividad": 11.27,
        "Fecha Inicio": 13.73,
        "Fecha Fin": 13.0,
        "Duración (min)": 14.0,
        "Completado": 8.0,
        "Planta": 7.18,
        "Tipo Actividad": 8.91,
        "Importancia": 12.54,
        "Tiempo a Invertir": 9.27,
        "Notas / Historial": 30.91
    }
    
    def format_sheet(ws, fill_header, widths_dict, is_todo_sheet=False):
        if not is_todo_sheet:
            ws.sheet_view.zoomScale = 90
            
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = font_bold
            cell.fill = fill_header
            header_val = str(cell.value or "")
            
            if "id" in header_val.lower() and not is_todo_sheet:
                cell.alignment = align_header_left
            else:
                if any(x in header_val.lower() for x in ["título", "titulo", "descrip", "nota"]):
                    cell.alignment = align_header_center
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    
        if not is_todo_sheet:
            for row in range(2, ws.max_row + 1):
                for col_idx in range(1, ws.max_column + 1):
                    cell = ws.cell(row=row, column=col_idx)
                    cell.font = font_data
                    header_val = str(ws.cell(row=1, column=col_idx).value or "")
                    
                    if "id" in header_val.lower():
                        cell.alignment = align_data_center
                    elif any(x in header_val.lower() for x in ["título", "titulo", "descrip", "contacto", "nota"]):
                        cell.alignment = align_data_left_wrap
                    else:
                        cell.alignment = align_data_left
                        
        for col_idx in range(1, ws.max_column + 1):
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            header_val = str(ws.cell(row=1, column=col_idx).value or "")
            if header_val in widths_dict:
                ws.column_dimensions[col_letter].width = widths_dict[header_val]
            else:
                ws.column_dimensions[col_letter].width = 15

    # 1. Hoja de Consolidado General (Trabajo, Gestión y Actividades todo junto)
    ws_consolidado = wb.active
    ws_consolidado.title = "Consolidado"
    headers_consolidado = [
        "ID Trabajo", "Título Trabajo", "Estado Trabajo", "Contacto/Cliente",
        "ID Gestión", "Título Gestión", "Estado Gestión", "Prioridad Gestión",
        "ID Actividad", "Título Actividad", "Descripción Actividad", "Estado Actividad",
        "Fecha Inicio", "Fecha Fin", "Duración (min)", "Completado",
        "Planta", "Tipo Actividad", "Importancia", "Tiempo a Invertir", "Notas / Historial"
    ]
    ws_consolidado.append(headers_consolidado)
    
    actividades = db.exec(select(Actividad).order_by(Actividad.id)).all()
    actividades_listadas = set()
    tareas_listadas = set()
    trabajos_listados = set()
    
    # Agregar actividades y sus ancestros
    for act in actividades:
        t_id = act.trabajo_id
        g_id = act.tarea_id
        job = act.trabajo
        tarea = act.tarea
        
        ws_consolidado.append([
            t_id or "",
            job.titulo if job else "Independiente",
            job.estado if job else "",
            job.cliente_contacto if job else "",
            g_id or "",
            tarea.titulo if tarea else "Independiente",
            tarea.estado if tarea else "",
            tarea.prioridad if tarea else "",
            act.id,
            act.titulo,
            act.descripcion or "",
            act.estado,
            act.fecha_inicio.strftime("%d-%m-%y %H:%M") if act.fecha_inicio else "",
            act.fecha_fin.strftime("%d-%m-%y %H:%M") if act.fecha_fin else "",
            act.duracion_minutos or 0,
            "SÍ" if act.completado else "NO",
            act.planta or "",
            act.tipo_actividad or "",
            act.importancia or "",
            act.tiempo_invertir or "",
            act.nota or ""
        ])
        actividades_listadas.add(act.id)
        if g_id:
            tareas_listadas.add(g_id)
        if t_id:
            trabajos_listados.add(t_id)
            
    # Agregar Gestiones que no tienen actividades
    tareas = db.exec(select(Tarea).order_by(Tarea.id)).all()
    for tr in tareas:
        if tr.id not in tareas_listadas:
            job = tr.trabajo
            ws_consolidado.append([
                tr.trabajo_id or "",
                job.titulo if job else "Independiente",
                job.estado if job else "",
                job.cliente_contacto if job else "",
                tr.id,
                tr.titulo,
                tr.estado,
                tr.prioridad,
                "", "", "", "", "", "", "", "",
                tr.planta or "",
                tr.tipo_actividad or "",
                tr.importancia or "",
                "",
                tr.nota or ""
            ])
            tareas_listadas.add(tr.id)
            if tr.trabajo_id:
                trabajos_listados.add(tr.trabajo_id)
                
    # Agregar Trabajos que no tienen actividades ni gestiones
    trabajos = db.exec(select(Trabajo).order_by(Trabajo.id)).all()
    for t in trabajos:
        if t.id not in trabajos_listados:
            ws_consolidado.append([
                t.id,
                t.titulo,
                t.estado,
                t.cliente_contacto or "",
                "", "", "", "", "", "", "", "", "", "", "", "",
                t.planta or "",
                t.tipo_actividad or "",
                t.importancia or "",
                "",
                t.nota or ""
            ])
            trabajos_listados.add(t.id)
            
    fill_header_consolidado = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    format_sheet(ws_consolidado, fill_header_consolidado, widths_consolidado)
    ws_consolidado.sheet_view.zoomScale = 80

    # 2. Hoja de Trabajos
    ws_trabajos = wb.create_sheet(title="Trabajos")
    headers_trabajos = ["ID Trabajo", "Título Trabajo", "Descripción", "Contacto/Cliente", "Estado", "Color", "Fecha Creación", "Nota", "Planta", "Tipo Actividad", "Importancia"]
    ws_trabajos.append(headers_trabajos)
    
    trabajos = db.exec(select(Trabajo).order_by(Trabajo.orden)).all()
    for t in trabajos:
        ws_trabajos.append([
            t.id,
            t.titulo,
            t.descripcion or "",
            t.cliente_contacto or "",
            t.estado,
            t.color or "",
            t.fecha_creacion.strftime("%d-%m-%y %H:%M") if t.fecha_creacion else "",
            t.nota or "",
            t.planta or "",
            t.tipo_actividad or "",
            t.importancia or ""
        ])
    
    fill_header_trabajos = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    format_sheet(ws_trabajos, fill_header_trabajos, widths_trabajos)

    # 2. Hoja de Gestiones
    ws_tareas = wb.create_sheet(title="Gestiones")
    headers_tareas = ["ID Trabajo", "Título Trabajo", "ID Gestión", "Título Gestión", "Descripción", "Prioridad", "Estado", "Fecha Creación", "Nota", "Planta", "Tipo Actividad", "Importancia"]
    ws_tareas.append(headers_tareas)
    
    tareas = db.exec(select(Tarea).order_by(Tarea.trabajo_id, Tarea.orden)).all()
    for tr in tareas:
        titulo_trabajo = tr.trabajo.titulo if tr.trabajo else "Independiente"
        ws_tareas.append([
            tr.trabajo_id or "",
            titulo_trabajo,
            tr.id,
            tr.titulo,
            tr.descripcion or "",
            tr.prioridad,
            tr.estado,
            tr.fecha_creacion.strftime("%d-%m-%y %H:%M") if tr.fecha_creacion else "",
            tr.nota or "",
            tr.planta or "",
            tr.tipo_actividad or "",
            tr.importancia or ""
        ])
        
    fill_header_tareas = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
    format_sheet(ws_tareas, fill_header_tareas, widths_gestiones)

    # 3. Hoja de Actividades
    ws_actividades = wb.create_sheet(title="Actividades")
    headers_actividades = ["ID Actividad", "Título Actividad", "Descripción", "Fecha Inicio", "Fecha Fin", "Estado", "Color", "ID Trabajo", "Título Trabajo", "ID Gestión", "Título Gestión", "Fecha Creación", "Nota", "Planta", "Tipo Actividad", "Importancia", "Tiempo a Invertir", "Completado"]
    ws_actividades.append(headers_actividades)
    
    actividades = db.exec(select(Actividad).order_by(Actividad.orden)).all()
    for act in actividades:
        titulo_trabajo = act.trabajo.titulo if act.trabajo else "Independiente"
        titulo_tarea = act.tarea.titulo if act.tarea else "Independiente"
        ws_actividades.append([
            act.id,
            act.titulo,
            act.descripcion or "",
            act.fecha_inicio.strftime("%d-%m-%y %H:%M") if act.fecha_inicio else "",
            act.fecha_fin.strftime("%d-%m-%y %H:%M") if act.fecha_fin else "",
            act.estado,
            act.color or "",
            act.trabajo_id or "",
            titulo_trabajo,
            act.tarea_id or "",
            titulo_tarea,
            act.fecha_creacion.strftime("%d-%m-%y %H:%M") if act.fecha_creacion else "",
            act.nota or "",
            act.planta or "",
            act.tipo_actividad or "",
            act.importancia or "",
            act.tiempo_invertir or "",
            "SÍ" if act.completado else "NO"
        ])
        
    fill_header_actividades = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
    format_sheet(ws_actividades, fill_header_actividades, widths_actividades)

    # 4. Hoja de To-Dos rápidos (ahora Actividades Rápidas)
    ws_todos = wb.create_sheet(title="Lista To-Do Rápida")
    headers_todos = ["ID Actividad", "Título Actividad", "Completado", "Fecha Creación", "Descripción", "Planta", "Importancia"]
    ws_todos.append(headers_todos)
    
    todos = db.exec(select(Actividad).where(Actividad.trabajo_id == None).where(Actividad.tarea_id == None).order_by(Actividad.orden)).all()
    for td in todos:
        ws_todos.append([
            td.id,
            td.titulo,
            "SÍ" if td.completado else "NO",
            td.fecha_creacion.strftime("%d-%m-%y %H:%M") if td.fecha_creacion else "",
            td.descripcion or "",
            td.planta or "",
            td.importancia or ""
        ])
        
    fill_header_todos = PatternFill(start_color="6B7280", end_color="6B7280", fill_type="solid")
    format_sheet(ws_todos, fill_header_todos, widths_todos, is_todo_sheet=True)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    headers = {
        'Content-Disposition': f'attachment; filename="agenda_reporte_{datetime.now().strftime("%d-%m-%y_%H%M%S")}.xlsx"'
    }
    return StreamingResponse(stream, headers=headers, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# Montar archivos estáticos para la interfaz
app.mount("/static", StaticFiles(directory="static"), name="static")
